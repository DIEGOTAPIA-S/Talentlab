import os
import shutil
import pptx

src = "PITCH PRESENTACIÓN FINAL.pptx"
backup = "PITCH PRESENTACIÓN FINAL_PRE_FILL_BACKUP.pptx"
if not os.path.exists(backup):
    shutil.copy2(src, backup)
    print("Backup created:", backup)

prs = pptx.Presentation(src)
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            text = tf.text.strip()
            if text:
                print(f"-- Shape {s_idx} ({shape.name}) --")
                for p in tf.paragraphs:
                    if p.text.strip():
                        print(f"   [P]: {p.text.strip()}")
        elif shape.has_table:
            print(f"-- Table {s_idx} ({len(shape.table.rows)}x{len(shape.table.columns)}) --")
            for r in shape.table.rows:
                row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
                print(f"   {row_txt}")
