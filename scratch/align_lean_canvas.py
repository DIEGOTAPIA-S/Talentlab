import os
import shutil
import docx
from docx.shared import Pt, RGBColor

# 1. Backups
docx_path = "MATRIZ LEAN CANVAS V1.docx"
docx_backup = "MATRIZ LEAN CANVAS V1_PRE_ALIGN_BACKUP.docx"
shutil.copy2(docx_path, docx_backup)

md_path = "Matriz_Lean_Canvas.md"
md_backup = "Matriz_Lean_Canvas_PRE_ALIGN_BACKUP.md"
shutil.copy2(md_path, md_backup)

# 2. Update DOCX
doc = docx.Document(docx_path)
table = doc.tables[0]

def set_cell_paragraphs(cell, items):
    # clear existing paragraphs except first
    p_first = cell.paragraphs[0]
    p_first.text = ""
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    
    for idx, (title, text, bold_prefix, italic) in enumerate(items):
        p = p_first if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        
        if title:
            r1 = p.add_run(title)
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            if not text:
                continue
        
        if bold_prefix:
            r_bp = p.add_run(bold_prefix)
            r_bp.font.size = Pt(9)
            r_bp.font.bold = True
            
        if text:
            r2 = p.add_run(text)
            r2.font.size = Pt(9)
            r2.font.bold = False
            r2.font.italic = bool(italic)

# Cell (0, 4) - Ventaja especial
ventaja_items = [
    ("Ventaja especial:", "", False, False),
    ("", " Modelos de Machine Learning optimizados y entrenados con datos históricos de deserción contextualizados a la realidad socioeconómica colombiana (frente a plataformas norteamericanas no adaptadas).", "Algoritmos Predictivos Propios:", False),
    ("", " Estructura multi-tenant altamente eficiente que reduce el costo de servidor por estudiante activo a menos de $0.05 USD anuales, lo que permite una flexibilidad de precios invencible.", "Arquitectura Serverless Optimizada (Bajo Costo):", False),
    ("", " Respaldo formativo y validación técnica directa por parte del equipo de desarrollo del programa Tecnólogo en Análisis y Desarrollo de Software del SENA.", "Validación Técnica ADSO - SENA:", False),
]
set_cell_paragraphs(table.cell(0, 4), ventaja_items)

# Cell (2, 0) - Estructura de Costes
costes_items = [
    ("Estructura de Costes:", "", False, False),
    ("", " $14.350.000 COP ($9.850.000 en 3 portátiles de alto rendimiento para desarrollo de software y conectividad; $4.500.000 en estaciones ergonómicas de trabajo y mobiliario).", "Costos de Equipos y Dotación Inicial (Activos Fijos):", False),
    ("", " Constitución legal de la S.A.S. en Cámara de Comercio, registro de marca ante la SIC e inscripción del software: $2.804.000 COP.", "Gastos Preoperativos y Legales (Activos Intangibles):", False),
    ("", " $30.190.222 COP (Fondos de maniobra para nómina técnica, servidores e imprevistos durante la fase de puesta en marcha).", "Capital de Trabajo Operativo Inicial:", False),
    ("", " Servidores en la nube (PostgreSQL/Supabase, AWS Lambda Serverless, APIs NLP): ~$300.000 COP/mes ($3.600.000 COP/año). Costo variable unitario por licencia: $312.000 COP (Margen Bruto del 93,07%).", "Costos Operativos Cloud y Variables:", False),
    ("", " $47.344.222 COP.", "Total Inversión Inicial Requerida:", True),
    ("", " Financiado mediante Crédito Bancario Comercial por $40.644.222 COP (85,85% a 5 años, tasa 14% E.A.) y Aporte Propio de Socios por $6.700.000 COP (14,15%). Fondo Emprender como opción alternativa complementaria.", "Estructura de Financiación:", False),
]
set_cell_paragraphs(table.cell(2, 0), costes_items)

# Cell (2, 3) - Fuentes de ingreso
ingresos_items = [
    ("Fuentes de ingreso:", "", False, False),
    ("Modelo B2B SaaS de Suscripción Anual (Premium Multi-tenant):", "", True, False),
    ("", " $4.500.000 COP / año (Para instituciones técnicas y tecnológicas con hasta 1.000 estudiantes; alertas de riesgo tempranas y simulador base).", "• Licencia Plan Lean (Básica):", False),
    ("", " $6.500.000 COP / año (Hasta 2.500 estudiantes; alertas avanzadas, simulador NLP completo con feedback instantáneo y red de mentorías).", "• Licencia Plan Growth (Institucional):", False),
    ("", " $9.000.000 COP / año (Más de 2.500 estudiantes; parametrización personalizada, integración API dedicada y soporte 24/7).", "• Licencia Plan Enterprise:", False),
    ("", " Carga inicial de bases de datos históricas, integración API con SIS/LMS institucional (Q10, Moodle, Banner) y capacitación docente.", "Servicios Profesionales de Consultoría e Integración (CIIU 6202):", False),
    ("", " Año 1: 12 licencias ($54.000.000 COP) escalando progresivamente hasta Año 5: 68 licencias ($358.331.000 COP).", "Proyección de Ventas a 5 Años:", False),
    ("", " Punto de Equilibrio: 32,2 licencias Plan Lean | VPN (TIO 15%): +$125.340.830 COP | TIR: 62,88% | PRI: 2,4 años | Relación B/C: 1,30.", "Viabilidad Financiera:", True),
]
set_cell_paragraphs(table.cell(2, 3), ingresos_items)

doc.save(docx_path)
print("Updated MATRIZ LEAN CANVAS V1.docx successfully!")

# 3. Update MD file
md_content = """# MATRIZ LEAN CANVAS – TALENTLAB

El modelo Lean Canvas permite sintetizar la estrategia comercial y técnica del proyecto productivo **TalentLab**, enfocándose en el encaje producto-mercado (*Product-Market Fit*) para una plataforma de software como servicio (**SaaS B2B Multi-tenant**).

---

## Representación Visual del Lean Canvas

| **1. PROBLEMA** | **2. SOLUCIÓN** | **4. PROPUESTA DE VALOR ÚNICA** | **5. VENTAJA INJUSTA** | **3. SEGMENTOS DE CLIENTES** |
| :--- | :--- | :--- | :--- | :--- |
| **Top 3 Problemas:**<br>1. **Altas tasas de deserción académica** (supera el 33% en educación técnica/tecnológica en Colombia según el MEN/SPADIES) debido a factores socioeconómicos y académicos no detectados a tiempo.<br>2. **Brecha de empleabilidad laboral ("Job Readiness")**: Los estudiantes carecen de habilidades blandas y preparación real para entrevistas técnicas/comportamentales.<br>3. **Desconexión con Egresados**: Falta de canales dinámicos y estructurados para que los graduados guíen a estudiantes activos (capital social desaprovechado).<br><br>**Alternativas Existentes:**<br>- Sistemas LMS/SIS tradicionales (Moodle, Q10, Banner) sin módulos de analítica predictiva.<br>- Simuladores de entrevista genéricos en inglés y de alto costo corporativo.<br>- Grupos informales en redes sociales (Facebook/LinkedIn) sin vinculación institucional. | **Top 3 Características:**<br>1. **Semáforo Predictivo de Deserción:** Panel administrativo impulsado por Machine Learning (algoritmo *Random Forest*) que analiza variables académicas, socioeconómicas e institucionales en tiempo real.<br>2. **Simulador de Entrevistas con IA:** Motor conversacional de procesamiento de lenguaje natural (NLP) que simula entrevistas técnicas/blandas en español con retroalimentación automatizada instantánea.<br>3. **Red Multi-tenant de Mentorías:** Módulo social privado para la interacción segura entre estudiantes activos y egresados mediante *matching* algorítmico de perfiles formativos. | **Mensaje Claro y Convincente:**<br>### *"Detecta. Prepara. Conecta."*<br><br>Una plataforma predictiva e interactiva de acompañamiento estudiantil integral que reduce la deserción académica en un 15% promedio y entrena a los estudiantes con IA para asegurar su éxito en el mercado laboral real, centralizando todo el ciclo del estudiante en un entorno unificado y de bajo costo operativo.<br><br>**Concepto de Alto Nivel:**<br>*"El 'Salesforce' de la retención y éxito estudiantil para instituciones de educación técnica y superior en Latinoamérica."* | **Lo que no se puede copiar/comprar fácilmente:**<br>1. **Algoritmos Predictivos Propios:** Modelos de Machine Learning optimizados y entrenados con datos históricos de deserción contextualizados a la realidad socioeconómica colombiana (frente a plataformas norteamericanas no adaptadas).<br>2. **Arquitectura Serverless Optimizada (Bajo Costo):** Estructura multi-tenant altamente eficiente que reduce el costo de servidor por estudiante activo a menos de **$0.05 USD anuales**, lo que permite una flexibilidad de precios invencible.<br>3. **Respaldo y Validación ADSO - SENA:** Validación técnica directa por parte del equipo formativo del Tecnólogo en Análisis y Desarrollo de Software del SENA. | **Segmento Meta:**<br>1. **B2B (Mercado Principal):** Directores de bienestar institucional, decanos académicos y rectores de instituciones privadas de educación técnica, tecnológica e instituciones universitarias de nivel intermedio en la región andina colombiana (Mercado SOM: 80 instituciones).<br>2. **B2C (Mercado Secundario):** Estudiantes y egresados que acceden a la plataforma a través de las licencias institucionales.<br><br>**Early Adopters:**<br>- Instituciones educativas técnicas y tecnológicas privadas en Bogotá y Cundinamarca que cuenten con altos índices de deserción en los primeros tres semestres académicos y busquen acreditar sus procesos de calidad. |
| **9. MÉTRICAS CLAVE** | | | | **6. CANALES** |
| *   **Tasa de Deserción Anual:** Reducción porcentual de la deserción académica institucional (Meta: Mitigación del 15% promedio).<br>*   **Engagement del Estudiante:** Frecuencia de uso del simulador de entrevistas con IA y participación activa en el módulo de mentorías.<br>*   **Churn Rate de Clientes B2B:** Porcentaje de desuscripción de instituciones educativas premium (Meta: Retención anual > 92%).<br>*   **NPS (Net Promoter Score):** Nivel de satisfacción de directivos de bienestar y estudiantes activos (Meta: > 75 puntos). | | | | **Caminos hacia el Cliente:**<br>1. **Venta Directa B2B (Consultoría Informática):** Reuniones comerciales estructuradas con directores de bienestar y rectores educativos apoyadas por demostraciones técnicas del MVP.<br>2. **Inbound Marketing & SEO:** Creación de contenido especializado (Whitepapers y artículos de blog) sobre retención de estudiantes y analítica predictiva en la educación en LATAM (CIIU 6202).<br>3. **Eventos y Congresos de EdTech:** Participación activa en foros de innovación educativa de la Red de Instituciones de Educación Superior y el MEN.<br>4. **Canal Social Freemium:** Acceso gratuito limitado para escuelas públicas de escasos recursos, actuando como motor orgánico de visibilidad nacional. |

---

| **7. ESTRUCTURA DE COSTES** | **8. FLUJOS DE INGRESOS** |
| :--- | :--- |
| **Costos de Equipos y Dotación Inicial (Activos Fijos):**<br>- $14.350.000 COP ($9.850.000 en 3 portátiles de alto rendimiento para desarrollo de software y conectividad; $4.500.000 en estaciones ergonómicas de trabajo y mobiliario).<br>**Gastos Preoperativos y Legales (Activos Intangibles):**<br>- Constitución legal S.A.S. en Cámara de Comercio, registro marcario ante la SIC e inscripción del software: $2.804.000 COP.<br>**Capital de Trabajo Operativo Inicial:**<br>- $30.190.222 COP (Fondos de maniobra para nómina técnica, servidores e imprevistos durante la fase de puesta en marcha).<br>**Costos Operativos Cloud y Variables:**<br>- Servidores en la nube (PostgreSQL/Supabase, AWS Lambda Serverless, APIs NLP): ~$300.000 COP/mes ($3.600.000 COP/año). Costo variable unitario por licencia: $312.000 COP (Margen Bruto del 93,07%).<br>**Total Inversión Inicial Requerida:**<br>- **$47.344.222 COP**.<br>**Estructura de Financiación:**<br>- **Crédito Bancario Comercial:** **$40.644.222 COP (85,85%)** a 5 años (60 meses) y tasa del **14% E.A.**<br>- **Aporte Propio de Socios:** **$6.700.000 COP (14,15%)**.<br>*(Fondo Emprender como opción alternativa complementaria).* | **Modelo B2B SaaS de Suscripción Anual (Premium Multi-tenant):**<br>- **Plan Lean (Básico):** **$4.500.000 COP / año** (Para instituciones de hasta 1.000 estudiantes; alertas de riesgo y simulador base).<br>- **Plan Growth (Institucional):** **$6.500.000 COP / año** (Hasta 2.500 estudiantes; alertas avanzadas, simulador NLP completo y red de mentorías).<br>- **Plan Enterprise:** **$9.000.000 COP / año** (Más de 2.500 estudiantes; parametrización personalizada, API dedicada y soporte prioritario).<br>**Servicios Profesionales de Consultoría e Integración (CIIU 6202):**<br>- Carga y migración inicial de bases de datos históricas, integración API con SIS/LMS institucional (Q10, Moodle, Banner) y capacitación docente.<br><br>**Proyección de Ventas a 5 Años:**<br>- **Año 1:** 12 licencias &rarr; **$54.000.000 COP**.<br>- **Año 2:** 24 licencias &rarr; **$114.156.000 COP**.<br>- **Año 3:** 38 licencias &rarr; **$189.691.000 COP**.<br>- **Año 4:** 52 licencias &rarr; **$267.094.000 COP**.<br>- **Año 5:** 68 licencias &rarr; **$358.331.000 COP**.<br><br>**Viabilidad Financiera Consolidada:**<br>- **Punto de Equilibrio:** 32,2 licencias Plan Lean ($144,9M COP).<br>- **VPN (TIO 15%):** **+$125.340.830 COP** | **TIR:** **62,88%**.<br>- **PRI:** **2,4 años** (mes 29) | **Relación B/C:** **1,30**. |

---

## Notas de Implementación y Encaje Estratégico

El diseño de este Lean Canvas se sustenta en una rigurosa arquitectura de costos controlada. La elección del **Stack Tecnológico** (Backend Django REST, Base de Datos Supabase/PostgreSQL y Frontend en React) responde directamente a la necesidad de mantener el costo operativo variable a niveles extremadamente competitivos. 

El modelo de predicción de deserción, construido sobre librerías de código abierto (*scikit-learn*), elimina la dependencia de costosos licenciamientos de software propietario de terceros, garantizando que el flujo de caja operativo de TalentLab soporte con solvencia el punto de equilibrio proyectado (32,2 licencias) y la generación neta de valor reflejada en un **VPN de +$125.340.830 COP** y una **TIR de 62,88%** con un **Periodo de Recuperación de Inversión de 2,4 años**.
"""

with open(md_path, "w", encoding="utf-8") as f:
    f.write(md_content)
print("Updated Matriz_Lean_Canvas.md successfully!")
