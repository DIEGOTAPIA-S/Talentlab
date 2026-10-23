import pptx

prs = pptx.Presentation("PITCH PRESENTACIÓN FINAL.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"=== SLIDE {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            first_line = shape.text_frame.text.split('\n')[0][:50]
            print(f"  Shape {s_idx} ({shape.name}): '{first_line}' [L:{shape.left}, T:{shape.top}, W:{shape.width}, H:{shape.height}]")
