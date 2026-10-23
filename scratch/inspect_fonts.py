import docx

doc = docx.Document("MATRIZ LEAN CANVAS V1.docx")
table = doc.tables[0]

for r in range(len(table.rows)):
    for c in range(len(table.columns)):
        cell = table.cell(r, c)
        for p in cell.paragraphs:
            for run in p.runs:
                if run.font.name or run.font.size:
                    print(f"Cell ({r},{c}) run: font={run.font.name}, size={run.font.size}")
                    break
            if p.runs and (p.runs[0].font.name or p.runs[0].font.size):
                break
