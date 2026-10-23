import pptx

prs = pptx.Presentation("PITCH PRESENTACIÓN FINAL.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        print(f"Shape {s_idx}: name='{shape.name}', type={shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                runs_info = []
                for r in p.runs:
                    runs_info.append(f"'{r.text}'(sz={r.font.size}, b={r.font.bold}, color={r.font.color.type if r.font.color else None})")
                print(f"   P{p_idx}: {' '.join(runs_info)}")
